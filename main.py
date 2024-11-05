import re
import argparse
from tabulate import tabulate
from colorama import Fore, Style, init
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
import xlsxwriter
import xlsxwriter
from xlsxwriter.utility import xl_col_to_name

init()


def find_active_cpu(file_path):
    """
    Find the active core by the number of clocks and instructions.
    """
    cpu_data = {}
    pattern_cycles = re.compile(r'system\.clusters\.(cpu\d+)\.numCycles\s+(\d+)')
    pattern_insts_committed = re.compile(r'system\.clusters\.(cpu\d+)\.committedInsts\s+(\d+)')

    with open(file_path, 'r') as file:
        for line in file:
            if match := pattern_cycles.search(line):
                cpu = match.group(1)
                cycles = int(match.group(2))
                if cpu not in cpu_data:
                    cpu_data[cpu] = {'cycles': cycles, 'committedInsts': 0}
                else:
                    cpu_data[cpu]['cycles'] = cycles
            if match := pattern_insts_committed.search(line):
                cpu = match.group(1)
                committed_insts = int(match.group(2))
                if cpu not in cpu_data:
                    cpu_data[cpu] = {'cycles': 0, 'committedInsts': committed_insts}
                else:
                    cpu_data[cpu]['committedInsts'] = committed_insts

    active_cpu = max(cpu_data, key=lambda x: (cpu_data[x]['cycles'], cpu_data[x]['committedInsts']))
    print(
        f"Active CPU is {active_cpu} with {cpu_data[active_cpu]['cycles']} cycles and {cpu_data[active_cpu]['committedInsts']} committed instructions")
    return active_cpu


def extract_cpu_statistics(file_path, active_cpu):
    stats = {}

    # CPU statistics
    pattern_cycles = re.compile(rf'system\.clusters\.{active_cpu}\.numCycles\s+(\d+)')
    pattern_insts_issued = re.compile(rf'system\.clusters\.{active_cpu}\.instsIssued\s+(\d+)')
    pattern_insts_committed = re.compile(rf'system\.clusters\.{active_cpu}\.committedInsts\s+(\d+)')
    pattern_cpi = re.compile(rf'system\.clusters\.{active_cpu}\.cpi\s+([\d\.]+)')
    pattern_ipc = re.compile(rf'system\.clusters\.{active_cpu}\.ipc\s+([\d\.]+)')

    # Load/Store Queue statistics
    pattern_lsq_forw_loads = re.compile(rf'system\.clusters\.{active_cpu}\.lsq\d+\.forwLoads\s+(\d+)')
    pattern_lsq_squashed_loads = re.compile(rf'system\.clusters\.{active_cpu}\.lsq\d+\.squashedLoads\s+(\d+)')
    pattern_lsq_squashed_stores = re.compile(rf'system\.clusters\.{active_cpu}\.lsq\d+\.squashedStores\s+(\d+)')
    pattern_lsq_ignored_responses = re.compile(rf'system\.clusters\.{active_cpu}\.lsq\d+\.ignoredResponses\s+(\d+)')
    pattern_lsq_mem_order_violations = re.compile(rf'system\.clusters\.{active_cpu}\.lsq\d+\.memOrderViolation\s+(\d+)')
    pattern_lsq_rescheduled_loads = re.compile(rf'system\.clusters\.{active_cpu}\.lsq\d+\.rescheduledLoads\s+(\d+)')
    pattern_lsq_blocked_by_cache = re.compile(rf'system\.clusters\.{active_cpu}\.lsq\d+\.blockedByCache\s+(\d+)')

    # Functional Unit usage patterns
    pattern_fu_busy = re.compile(rf'system\.clusters\.{active_cpu}\.statFuBusy::(\w+)\s+(\d+)\s+([\d\.]+)%')

    # Cache patterns
    pattern_cache_hits = re.compile(rf'system\.clusters\.{active_cpu}\.dcache\.overallHits::total\s+(\d+)')
    pattern_cache_misses = re.compile(rf'system\.clusters\.{active_cpu}\.dcache\.overallMisses::total\s+(\d+)')
    pattern_cache_miss_rate = re.compile(rf'system\.clusters\.{active_cpu}\.dcache\.overallMissRate::total\s+([\d\.]+)')
    pattern_cache_latency = re.compile(
        rf'system\.clusters\.{active_cpu}\.dcache\.overallAvgMissLatency::total\s+([\d\.]+)')

    # Memory Dependency Unit patterns
    pattern_mem_dep_unit = re.compile(rf'system\.clusters\.{active_cpu}\.MemDepUnit__(\d+)\.(\w+)\s+(\d+)')

    # Branch prediction patterns
    pattern_bp_btb_lookups = re.compile(rf'system\.clusters\.{active_cpu}\.branchPred\.BTBLookups\s+(\d+)')
    pattern_bp_btb_hits = re.compile(rf'system\.clusters\.{active_cpu}\.branchPred\.BTBHits\s+(\d+)')
    pattern_bp_btb_hit_ratio = re.compile(rf'system\.clusters\.{active_cpu}\.branchPred\.BTBHitRatio\s+([\d\.]+)')

    # Memory Controller patterns
    pattern_mem_bw_read = re.compile(r'system\.mem_ctrls(\d+)\.dram\.bwRead::total\s+(\d+)')
    pattern_mem_bw_write = re.compile(r'system\.mem_ctrls(\d+)\.dram\.bwWrite::total\s+(\d+)')
    pattern_mem_read_bursts = re.compile(r'system\.mem_ctrls(\d+)\.dram\.readBursts\s+(\d+)')
    pattern_mem_write_bursts = re.compile(r'system\.mem_ctrls(\d+)\.dram\.writeBursts\s+(\d+)')
    pattern_mem_queue_latency = re.compile(r'system\.mem_ctrls(\d+)\.dram\.avgQueueLatency\s+([\d\.]+)')
    pattern_mem_accesses = re.compile(r'system\.mem_ctrls(\d+)\.dram\.accesses::total\s+(\d+)')

    mem_ctrl_data = {}

    with open(file_path, 'r') as file:
        for line in file:
            # CPU statistics
            if match := pattern_cycles.search(line):
                stats['numCycles'] = int(match.group(1))
            if match := pattern_insts_issued.search(line):
                stats['instsIssued'] = int(match.group(1))
            if match := pattern_insts_committed.search(line):
                stats['committedInsts'] = int(match.group(1))
            if match := pattern_cpi.search(line):
                stats['cpi'] = float(match.group(1))
            if match := pattern_ipc.search(line):
                stats['ipc'] = float(match.group(1))

            # Load/Store Queue statistics
            if match := pattern_lsq_forw_loads.search(line):
                stats['lsq_forw_loads'] = int(match.group(1))
            if match := pattern_lsq_squashed_loads.search(line):
                stats['lsq_squashed_loads'] = int(match.group(1))
            if match := pattern_lsq_squashed_stores.search(line):
                stats['lsq_squashed_stores'] = int(match.group(1))
            if match := pattern_lsq_ignored_responses.search(line):
                stats['lsq_ignored_responses'] = int(match.group(1))
            if match := pattern_lsq_mem_order_violations.search(line):
                stats['lsq_mem_order_violations'] = int(match.group(1))
            if match := pattern_lsq_rescheduled_loads.search(line):
                stats['lsq_rescheduled_loads'] = int(match.group(1))
            if match := pattern_lsq_blocked_by_cache.search(line):
                stats['lsq_blocked_by_cache'] = int(match.group(1))

            # Functional Unit statistics
            if match := pattern_fu_busy.search(line):
                fu_type = match.group(1)
                fu_busy_count = int(match.group(2))
                fu_busy_rate = float(match.group(3))
                if 'FU_Busy' not in stats:
                    stats['FU_Busy'] = {}
                stats['FU_Busy'][fu_type] = {'count': fu_busy_count, 'rate': fu_busy_rate}

            # Cache statistics
            if match := pattern_cache_hits.search(line):
                stats['cache_hits'] = int(match.group(1))
            if match := pattern_cache_misses.search(line):
                stats['cache_misses'] = int(match.group(1))
            if match := pattern_cache_miss_rate.search(line):
                stats['cache_miss_rate'] = float(match.group(1))
            if match := pattern_cache_latency.search(line):
                stats['cache_miss_latency'] = float(match.group(1))

            # Memory Dependency Unit statistics
            if match := pattern_mem_dep_unit.search(line):
                unit_id = match.group(1)
                event_type = match.group(2)
                value = int(match.group(3))
                if 'MemDepUnit' not in stats:
                    stats['MemDepUnit'] = {}
                if unit_id not in stats['MemDepUnit']:
                    stats['MemDepUnit'][unit_id] = {}
                stats['MemDepUnit'][unit_id][event_type] = value

            # Branch prediction statistics
            if match := pattern_bp_btb_lookups.search(line):
                stats['btb_lookups'] = int(match.group(1))
            if match := pattern_bp_btb_hits.search(line):
                stats['btb_hits'] = int(match.group(1))
            if match := pattern_bp_btb_hit_ratio.search(line):
                stats['btb_hit_ratio'] = float(match.group(1))

            # Memory Controller statistics
            if match := pattern_mem_bw_read.search(line):
                mem_ctrl_id = match.group(1)
                bw_read = int(match.group(2))
                if mem_ctrl_id not in mem_ctrl_data:
                    mem_ctrl_data[mem_ctrl_id] = {}
                mem_ctrl_data[mem_ctrl_id]['bw_read'] = bw_read
            if match := pattern_mem_bw_write.search(line):
                mem_ctrl_id = match.group(1)
                bw_write = int(match.group(2))
                mem_ctrl_data[mem_ctrl_id]['bw_write'] = bw_write
            if match := pattern_mem_read_bursts.search(line):
                mem_ctrl_id = match.group(1)
                read_bursts = int(match.group(2))
                mem_ctrl_data[mem_ctrl_id]['read_bursts'] = read_bursts
            if match := pattern_mem_write_bursts.search(line):
                mem_ctrl_id = match.group(1)
                write_bursts = int(match.group(2))
                mem_ctrl_data[mem_ctrl_id]['write_bursts'] = write_bursts
            if match := pattern_mem_queue_latency.search(line):
                mem_ctrl_id = match.group(1)
                queue_latency = float(match.group(2))
                mem_ctrl_data[mem_ctrl_id]['queue_latency'] = queue_latency
            if match := pattern_mem_accesses.search(line):
                mem_ctrl_id = match.group(1)
                accesses = int(match.group(2))
                mem_ctrl_data[mem_ctrl_id]['accesses'] = accesses

    # Calculate total bandwidth and balance
    total_read_bw = sum(mem_ctrl['bw_read'] for mem_ctrl in mem_ctrl_data.values())
    total_write_bw = sum(mem_ctrl['bw_write'] for mem_ctrl in mem_ctrl_data.values())
    for mem_ctrl_id, mem_ctrl in mem_ctrl_data.items():
        read_share = (mem_ctrl['bw_read'] / total_read_bw * 100) if total_read_bw else 0
        write_share = (mem_ctrl['bw_write'] / total_write_bw * 100) if total_write_bw else 0
        mem_ctrl['read_share'] = read_share
        mem_ctrl['write_share'] = write_share

    stats['mem_ctrl_data'] = mem_ctrl_data
    return stats


def filter_non_zero_rows(rows):
    """
    Filters out rows where all values (except the first label column) are zero.
    """
    return [row for row in rows if any(value != 'N/A' and value != 0 for value in row[1:])]

def save_to_excel(stats_list, file_labels, category, mem_ctrl_table, rows_ipc_cpi=None, headers_general=None,
                  headers_ipc_cpi=None):
    file_name = f"{category}_statistics.xlsx"
    workbook = xlsxwriter.Workbook(file_name)

    worksheet_general = workbook.add_worksheet(f"{category}_general")
    worksheet_general.write_row(0, 0, headers_general)

    for row_idx, row in enumerate(mem_ctrl_table, start=1):
        worksheet_general.write_row(row_idx, 0, row)

    if category == "cpu":
        worksheet_general.write_row(0, 0, headers_general)

        for row_idx, row in enumerate(mem_ctrl_table, start=1):
            worksheet_general.write_row(row_idx, 0, row)

        general_chart = workbook.add_chart({'type': 'column'})
        for i, file_label in enumerate(file_labels):
            general_chart.add_series({
                'name': file_label,
                'categories': f"'{category}_general'!A2:A{len(mem_ctrl_table) + 1}",
                'values': f"'{category}_general'!{xl_col_to_name(i + 1)}2:" \
                          f"{xl_col_to_name(i + 1)}{len(mem_ctrl_table) + 1}",
            })

        general_chart.set_title({'name': 'CPU General Statistics'})
        general_chart.set_x_axis({'name': 'Metrics'})
        general_chart.set_y_axis({'name': 'Value'})
        worksheet_general.insert_chart('E2', general_chart)

        worksheet_ipc_cpi = workbook.add_worksheet(f"{category}_ipc_cpi")
        worksheet_ipc_cpi.write_row(0, 0, headers_ipc_cpi)

        for row_idx, row in enumerate(rows_ipc_cpi, start=1):
            worksheet_ipc_cpi.write_row(row_idx, 0, row)

        ipc_cpi_chart = workbook.add_chart({'type': 'column'})
        for i, file_label in enumerate(file_labels):
            ipc_cpi_chart.add_series({
                'name': file_label,
                'categories': f"'{category}_ipc_cpi'!A2:A{len(rows_ipc_cpi) + 1}",
                'values': f"'{category}_ipc_cpi'!{xl_col_to_name(i + 1)}2:" \
                          f"{xl_col_to_name(i + 1)}{len(rows_ipc_cpi) + 1}",
            })

        ipc_cpi_chart.set_title({'name': 'CPU IPC and CPI Statistics'})
        ipc_cpi_chart.set_x_axis({'name': 'Metrics (CPI, IPC)'})
        ipc_cpi_chart.set_y_axis({'name': 'Value'})
        worksheet_ipc_cpi.insert_chart('E2', ipc_cpi_chart)

    elif category == "lsq":
        for row_idx, row in enumerate(mem_ctrl_table, start=1):
            worksheet_general.write_row(row_idx, 0, row)

            chart = workbook.add_chart({'type': 'column'})
            for i, file_label in enumerate(file_labels):
                chart.add_series({
                    'name': file_label,
                    'categories': f"'{category}_general'!A{row_idx + 1}",
                    'values': f"'{category}_general'!{xl_col_to_name(i + 1)}{row_idx + 1}",
                })
            chart.set_title({'name': f'{row[0]} Statistics'})
            chart.set_x_axis({'name': 'Files'})
            chart.set_y_axis({'name': 'Value'})
            worksheet_general.insert_chart(f"E{row_idx * 10}", chart)

    elif category == "fu":
        for row_idx, row in enumerate(mem_ctrl_table, start=1):
            worksheet_general.write_row(row_idx, 0, row)

        chart = workbook.add_chart({'type': 'column'})
        for i, file_label in enumerate(file_labels):
            chart.add_series({
                'name': file_label,
                'categories': f"'{category}_general'!A2:A{len(mem_ctrl_table) + 1}",
                'values': f"'{category}_general'!{xl_col_to_name(i + 1)}2:" \
                          f"{xl_col_to_name(i + 1)}{len(mem_ctrl_table) + 1}",
            })

        chart.set_title({'name': 'Functional Unit Usage'})
        chart.set_x_axis({'name': 'Functional Units'})
        chart.set_y_axis({'name': 'Usage Count'})
        worksheet_general.insert_chart('E2', chart)

    elif category == "cache":
        for row_idx, row in enumerate(mem_ctrl_table, start=1):
            worksheet_general.write_row(row_idx, 0, row)

            chart = workbook.add_chart({'type': 'column'})
            for i, file_label in enumerate(file_labels):
                chart.add_series({
                    'name': file_label,
                    'categories': f"'{category}_general'!A{row_idx + 1}",
                    'values': f"'{category}_general'!{xl_col_to_name(i + 1)}{row_idx + 1}",
                })
            chart.set_title({'name': f'{row[0]} Cache Statistics'})
            chart.set_x_axis({'name': 'Files'})
            chart.set_y_axis({'name': 'Value'})
            worksheet_general.insert_chart(f"E{row_idx * 10}", chart)

    elif category == "bp":
        for row_idx, row in enumerate(mem_ctrl_table, start=1):
            worksheet_general.write_row(row_idx, 0, row)

        chart = workbook.add_chart({'type': 'column'})
        for i, file_label in enumerate(file_labels):
            chart.add_series({
                'name': file_label,
                'categories': f"'{category}_general'!A2:A{len(mem_ctrl_table) + 1}",
                'values': f"'{category}_general'!{xl_col_to_name(i + 1)}2:" \
                          f"{xl_col_to_name(i + 1)}{len(mem_ctrl_table) + 1}",
            })

        chart.set_title({'name': 'Branch Prediction Statistics'})
        chart.set_x_axis({'name': 'Metrics'})
        chart.set_y_axis({'name': 'Value'})
        worksheet_general.insert_chart('E2', chart)

    elif category == "mem_ctrl":
        num_metrics = 4
        for idx, row_data in enumerate(mem_ctrl_table, start=1):
            if row_data[0].startswith("Memory Controller"):
                chart = workbook.add_chart({'type': 'column'})
                mem_ctrl_name = row_data[0]

                for metric_idx in range(num_metrics):
                    for file_idx, file_label in enumerate(file_labels):
                        col = 1 + (file_idx * num_metrics) + metric_idx
                        metric_name = headers_general[col]
                        chart.add_series({
                            'name': f"{file_label} {metric_name}",
                            'categories': [f"{category}_general", idx, 0, idx, 0],
                            'values': [f"{category}_general", idx, col, idx, col],
                        })

                chart.set_title({'name': f"{mem_ctrl_name} Statistics"})
                chart.set_x_axis({'name': 'Metric'})
                chart.set_y_axis({'name': 'Value'})
                chart.set_legend({'position': 'right'})

                worksheet_general.insert_chart(f"E{idx * 10}", chart)
    #TODO: add piecharts.
    elif category == "mem_ctrl_balance":
        try:
            read_share_col = headers_general.index("Read % Share")
            write_share_col = headers_general.index("Write % Share")
            print(f"Read % Share Column Index: {read_share_col}, Write % Share Column Index: {write_share_col}")
        except ValueError as e:
            print("Error finding columns for 'Read % Share' or 'Write % Share':", e)
            workbook.close()
            return

        for file_idx, file_label in enumerate(file_labels):
            read_pie_chart = workbook.add_chart({'type': 'pie'})
            read_pie_chart.set_title({'name': f"{file_label} Read Share Distribution"})
            read_pie_chart.add_series({
                'categories': f"'{category}_general'!A2:A{len(mem_ctrl_table) + 1}",
                'values': f"'{category}_general'!{xl_col_to_name(read_share_col)}2:" \
                          f"{xl_col_to_name(read_share_col)}{len(mem_ctrl_table) + 1}",
                'data_labels': {'percentage': True}
            })
            worksheet_general.insert_chart(f"E{file_idx * 20 + 1}", read_pie_chart)

            write_pie_chart = workbook.add_chart({'type': 'pie'})
            write_pie_chart.set_title({'name': f"{file_label} Write Share Distribution"})
            write_pie_chart.add_series({
                'categories': f"'{category}_general'!A2:A{len(mem_ctrl_table) + 1}",
                'values': f"'{category}_general'!{xl_col_to_name(write_share_col)}2:" \
                          f"{xl_col_to_name(write_share_col)}{len(mem_ctrl_table) + 1}",
                'data_labels': {'percentage': True}
            })
            worksheet_general.insert_chart(f"E{(file_idx * 20) + 15}", write_pie_chart)

    workbook.close()
    print(f"Excel file saved as {file_name}")



def display_statistics(stats_list, file_labels, category):
    headers = ["Metric"] + file_labels
    rows = []

    if category == "cpu":
        rows_general = [
            ["Cycles"] + [stats.get('numCycles', 'N/A') for stats in stats_list],
            ["Instructions Issued"] + [stats.get('instsIssued', 'N/A') for stats in stats_list],
            ["Instructions Committed"] + [stats.get('committedInsts', 'N/A') for stats in stats_list]
        ]

        rows_ipc_cpi = [
            ["CPI"] + [stats.get('cpi', 'N/A') for stats in stats_list],
            ["IPC"] + [stats.get('ipc', 'N/A') for stats in stats_list]
        ]

        headers_general = ["Metric"] + file_labels
        headers_ipc_cpi = ["Metric"] + file_labels
        print("\nGeneral CPU Statistics:")
        print(tabulate(rows_general, headers=headers_general, tablefmt="grid"))

        print("\nIPC and CPI Statistics:")
        print(tabulate(rows_ipc_cpi, headers=headers_ipc_cpi, tablefmt="grid"))

        save_to_excel(stats_list, file_labels, category, rows_general, rows_ipc_cpi, headers_general, headers_ipc_cpi)

    elif category == "lsq":
        rows = [
            ["Forwarded Loads"] + [stats.get('lsq_forw_loads', 'N/A') for stats in stats_list],
            ["Squashed Loads"] + [stats.get('lsq_squashed_loads', 'N/A') for stats in stats_list],
            ["Squashed Stores"] + [stats.get('lsq_squashed_stores', 'N/A') for stats in stats_list],
            ["Ignored Responses"] + [stats.get('lsq_ignored_responses', 'N/A') for stats in stats_list],
            ["Memory Order Violations"] + [stats.get('lsq_mem_order_violations', 'N/A') for stats in stats_list],
            ["Rescheduled Loads"] + [stats.get('lsq_rescheduled_loads', 'N/A') for stats in stats_list],
            ["Blocked By Cache"] + [stats.get('lsq_blocked_by_cache', 'N/A') for stats in stats_list],
        ]
        print(tabulate(rows, headers=headers, tablefmt="grid"))
        save_to_excel(stats_list, file_labels, category, rows, [], headers, [])


    elif category == "fu":

        headers_fu = ["Functional Unit"] + file_labels

        rows = []

        for fu_type in stats_list[0].get('FU_Busy', {}):
            row = [fu_type] + [stats.get('FU_Busy', {}).get(fu_type, {}).get('count', 0) for stats in stats_list]

            rows.append(row)

        rows = filter_non_zero_rows(rows)

        print(tabulate(rows, headers=headers_fu, tablefmt="grid"))

        save_to_excel(stats_list, file_labels, category, rows, [], headers_fu, [])


    elif category == "cache":
        rows = [
            ["Cache Hits"] + [stats.get('cache_hits', 'N/A') for stats in stats_list],
            ["Cache Misses"] + [stats.get('cache_misses', 'N/A') for stats in stats_list],
            ["Cache Miss Rate"] + [stats.get('cache_miss_rate', 'N/A') for stats in stats_list],
            ["Cache Miss Latency"] + [stats.get('cache_miss_latency', 'N/A') for stats in stats_list],
        ]
        print(tabulate(rows, headers=headers, tablefmt="grid"))
        save_to_excel(stats_list, file_labels, category, rows, [], headers, [])

    elif category == "bp":
        rows = [
            ["BTB Lookups"] + [stats.get('btb_lookups', 'N/A') for stats in stats_list],
            ["BTB Hits"] + [stats.get('btb_hits', 'N/A') for stats in stats_list],
            ["BTB Hit Ratio"] + [stats.get('btb_hit_ratio', 'N/A') for stats in stats_list],
        ]
        print(tabulate(rows, headers=headers, tablefmt="grid"))
        save_to_excel(stats_list, file_labels, category, rows, [], headers, [])


    elif category == "mem_ctrl":

        mem_ctrl_table = []

        headers_mem_ctrl = ["Memory Controller"]

        for file_label in file_labels:
            headers_mem_ctrl.extend(
                [f"{file_label} Read Bandwidth (Bytes/s)", f"{file_label} Write Bandwidth (Bytes/s)",

                 f"{file_label} Read Bursts", f"{file_label} Write Bursts"])

        mem_ctrl_ids = set()

        for stats in stats_list:
            mem_ctrl_ids.update(stats['mem_ctrl_data'].keys())

        for mem_ctrl_id in sorted(mem_ctrl_ids):

            row = [f"Memory Controller {mem_ctrl_id}"]

            for stats in stats_list:
                mem_ctrl = stats['mem_ctrl_data'].get(mem_ctrl_id, {})

                row.extend([

                    mem_ctrl.get('bw_read', 'N/A'),

                    mem_ctrl.get('bw_write', 'N/A'),

                    mem_ctrl.get('read_bursts', 'N/A'),

                    mem_ctrl.get('write_bursts', 'N/A')

                ])

            mem_ctrl_table.append(row)

        print(tabulate(mem_ctrl_table, headers=headers_mem_ctrl, tablefmt="grid"))

        save_to_excel(stats_list, file_labels, category, mem_ctrl_table, [], headers_mem_ctrl, [])


    elif category == "mem_ctrl_balance":
        mem_ctrl_balance_table = []
        for idx, stats in enumerate(stats_list, start=1):
            for mem_ctrl_id, mem_ctrl in stats['mem_ctrl_data'].items():
                mem_ctrl_balance_table.append([
                    f"Memory Controller {mem_ctrl_id}",
                    mem_ctrl.get('bw_read', 'N/A'),
                    mem_ctrl.get('bw_write', 'N/A'),
                    mem_ctrl.get('read_bursts', 'N/A'),
                    mem_ctrl.get('write_bursts', 'N/A'),
                    mem_ctrl.get('queue_latency', 'N/A'),
                    mem_ctrl.get('accesses', 'N/A'),
                    f"{mem_ctrl.get('read_share', 0):.2f}%",
                    f"{mem_ctrl.get('write_share', 0):.2f}%"
                ])
        headers_mem_ctrl_balance = ["Memory Controller", "Read Bandwidth (Bytes/s)", "Write Bandwidth (Bytes/s)",
                                    "Read Bursts",
                                    "Write Bursts", "Queue Latency", "Total Accesses", "Read % Share", "Write % Share"]
        print(tabulate(mem_ctrl_balance_table, headers=headers_mem_ctrl_balance, tablefmt="grid"))
        save_to_excel(stats_list, file_labels, category, mem_ctrl_balance_table, [], headers_mem_ctrl_balance, [])


def main():
    parser = argparse.ArgumentParser(description="Extract and display specific statistics from gem5 stat files.")
    parser.add_argument("file_paths", nargs='+', help="Paths to the stat files (up to 3)", metavar="FILE")
    parser.add_argument("--category", choices=["cpu", "lsq", "fu", "cache", "bp", "mem_ctrl", "mem_ctrl_balance"],
                        required=True, help="Category of statistics to display")

    args = parser.parse_args()

    if len(args.file_paths) > 3:
        print("Please provide up to 3 stat files for comparison.")
        return

    stats_list = []
    file_labels = []
    for file_path in args.file_paths:
        active_cpu = find_active_cpu(file_path)
        stats = extract_cpu_statistics(file_path, active_cpu)
        stats_list.append(stats)
        file_labels.append(file_path.split('/')[-1])

    display_statistics(stats_list, file_labels, args.category)


if __name__ == "__main__":
    main()
